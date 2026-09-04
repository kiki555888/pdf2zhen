#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
本地极致轻量解析器（纯文本翻译专用）
- 支持格式: .txt, .md, .epub, .docx
- 按页分割，每页生成一个 文档名_part_xxx/文档名_part_xxx.md 文件
- 输出结构: mineru_results/文档名/文档名_part_001/文档名_part_001.md
- 尽量保持原格式换行和段落
"""

import sys
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Optional

try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False


class LocalParser:
    def __init__(
        self,
        pages_per_file: int = 20,
        avg_chars_per_page: int = 500,
        chars_per_page: int = 0,
        lines_per_page: int = 0,
    ):
        self.chars_per_page = chars_per_page if chars_per_page > 0 else pages_per_file * avg_chars_per_page
        self.lines_per_page = lines_per_page
        if not HAS_DOCX:
            print("⚠️ python-docx 未安装，DOCX文件将被跳过。")
        print(f"📄 文本分割: {self.chars_per_page} 字符/页")

    def process_folder(self, input_folder: Path, output_root: Path) -> None:
        if not input_folder.is_dir():
            print(f"❌ 输入文件夹不存在: {input_folder}")
            return

        supported_extensions = {'.txt', '.md', '.text', '.epub', '.docx'}
        files = [f for f in input_folder.iterdir() if f.is_file() and f.suffix.lower() in supported_extensions]

        if not files:
            print("⚠️ 输入文件夹中未找到任何支持的纯文本文件（txt, md, epub, docx）。")
            return

        print(f"\n📂 发现 {len(files)} 个文件，开始处理...")
        print("📖 支持格式: TXT, MD, EPUB, DOCX（保持原格式）\n")

        total_success = 0
        total_failed = 0

        for file_path in files:
            try:
                doc_name = file_path.stem
                pages = self._parse_file(file_path)  # 返回页面列表
                if pages is None:
                    total_failed += 1
                    continue

                # 为每个页面生成 part 目录和 md 文件，格式为 文档名_part_001
                for idx, page_content in enumerate(pages, start=1):
                    part_dir_name = f"{doc_name}_part_{idx:03d}"
                    part_dir = output_root / doc_name / part_dir_name
                    part_dir.mkdir(parents=True, exist_ok=True)
                    md_file = part_dir / f"{part_dir_name}.md"
                    with open(md_file, "w", encoding="utf-8") as f:
                        f.write(page_content)
                    print(f"   ✅ 生成: {file_path.name} -> {md_file}")

                total_success += 1
            except Exception as e:
                print(f"   ❌ 处理失败 {file_path.name}: {e}")
                total_failed += 1

        print(f"\n🎉 本地解析完成！成功: {total_success}, 失败: {total_failed}")

    # ==================== 统一解析入口，返回页面列表 ====================
    def _parse_file(self, file_path: Path) -> Optional[List[str]]:
        suffix = file_path.suffix.lower()
        if suffix in (".txt", ".md", ".text"):
            return self._process_text(file_path)
        elif suffix == ".epub":
            return self._process_epub(file_path)
        elif suffix == ".docx":
            if HAS_DOCX:
                return self._process_docx(file_path)
            else:
                print(f"   ⏭️ 跳过DOCX（未安装python-docx）: {file_path.name}")
                return None
        else:
            print(f"   ⏭️ 跳过不支持的文件: {file_path.name}")
            return None

    # ==================== TXT / MD ====================
    def _process_text(self, file_path: Path) -> Optional[List[str]]:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
        except UnicodeDecodeError:
            try:
                with open(file_path, 'r', encoding='gbk') as f:
                    content = f.read()
            except Exception as e:
                print(f"      ❌ 无法读取文本文件: {e}")
                return None

        # 按行分割，保留空行
        lines = content.splitlines()
        cleaned_lines = []
        for line in lines:
            if line.strip() == "":
                cleaned_lines.append("")
            else:
                cleaned = re.sub(r' +', ' ', line.strip())
                cleaned_lines.append(cleaned)

        cleaned_text = "\n".join(cleaned_lines)
        pages = self._split_text_into_pages(cleaned_text)
        return pages

    # ==================== EPUB ====================
    def _process_epub(self, epub_path: Path) -> Optional[List[str]]:
        try:
            with zipfile.ZipFile(epub_path, 'r') as zip_ref:
                opf_path = None
                for name in zip_ref.namelist():
                    if name.endswith('content.opf') or name.endswith('package.opf'):
                        opf_path = name
                        break
                if not opf_path:
                    print(f"      ❌ 未找到 OPF 文件")
                    return None

                opf_content = zip_ref.read(opf_path).decode('utf-8')
                root = ET.fromstring(opf_content)
                ns = {'opf': 'http://www.idpf.org/2007/opf'}

                item_map = {}
                for item in root.findall('.//opf:item', ns):
                    item_id = item.get('id')
                    href = item.get('href')
                    media_type = item.get('media-type')
                    if href and media_type and 'xml' in media_type:
                        item_map[item_id] = href

                spine_items = []
                for itemref in root.findall('.//opf:itemref', ns):
                    ref_id = itemref.get('idref')
                    if ref_id in item_map:
                        spine_items.append(item_map[ref_id])

                all_text_parts = []
                for href in spine_items:
                    try:
                        content = zip_ref.read(href).decode('utf-8')
                    except KeyError:
                        opf_dir = Path(opf_path).parent
                        if str(opf_dir) != '.':
                            try:
                                content = zip_ref.read(str(opf_dir / href)).decode('utf-8')
                            except:
                                continue
                        else:
                            continue

                    content = re.sub(r'<br\s*/?>', '\n', content)
                    content = re.sub(r'</(p|div|h[1-6]|li|blockquote|pre)>', '\n', content)
                    content = re.sub(r'<[^>]+>', '', content)
                    content = re.sub(r' +', ' ', content)
                    all_text_parts.append(content.strip())

                full_text = "\n\n".join(all_text_parts)

                lines = full_text.splitlines()
                cleaned_lines = []
                for line in lines:
                    if line.strip() == "":
                        cleaned_lines.append("")
                    else:
                        cleaned = re.sub(r' +', ' ', line.strip())
                        cleaned_lines.append(cleaned)
                cleaned_text = "\n".join(cleaned_lines)
                pages = self._split_text_into_pages(cleaned_text)
                return pages
        except Exception as e:
            print(f"      ❌ EPUB解析失败: {e}")
            return None

    # ==================== DOCX ====================
    def _process_docx(self, docx_path: Path) -> Optional[List[str]]:
        try:
            doc = DocxDocument(docx_path)
        except Exception as e:
            print(f"      ❌ 读取DOCX失败: {e}")
            return None

        lines = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                cleaned = re.sub(r' +', ' ', text)
                lines.append(cleaned)
            else:
                lines.append("")

        for table in doc.tables:
            md_table = self._table_to_markdown(table)
            if md_table:
                lines.append("")
                lines.append(md_table)

        cleaned_text = "\n".join(lines)
        pages = self._split_text_into_pages(cleaned_text)
        return pages

    def _table_to_markdown(self, table) -> str:
        if not table.rows:
            return ""
        md_lines = []
        header_cells = [cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells]
        md_lines.append("| " + " | ".join(header_cells) + " |")
        md_lines.append("| " + " | ".join(["---"] * len(header_cells)) + " |")
        for row in table.rows[1:]:
            row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            while len(row_cells) < len(header_cells):
                row_cells.append("")
            md_lines.append("| " + " | ".join(row_cells) + " |")
        return "\n".join(md_lines)

    # ==================== 分页函数 ====================
    def _split_text_into_pages(self, text: str) -> List[str]:
        """
        按字符数分页，尽量在换行处切割。
        返回页面列表，每页为字符串。
        """
        if self.chars_per_page <= 0:
            return [text]

        pages = []
        start = 0
        length = len(text)
        while start < length:
            end = start + self.chars_per_page
            if end >= length:
                pages.append(text[start:])
                break
            seek = end
            while seek > start and text[seek] not in ('\n', '\r'):
                seek -= 1
            if seek == start:
                pages.append(text[start:end])
                start = end
            else:
                pages.append(text[start:seek])
                start = seek + 1
        return pages
